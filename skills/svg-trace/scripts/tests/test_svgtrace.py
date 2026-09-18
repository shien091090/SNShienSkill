import re
import sys
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import svgtrace as st  # noqa: E402


def test_no_subcommand_exits_2():
    with pytest.raises(SystemExit) as exc:
        st.main([])
    assert exc.value.code == 2


def test_error_is_reported_without_traceback(capsys, monkeypatch):
    def boom(args):
        raise st.SvgTraceError("測試用錯誤")

    monkeypatch.setattr(st, "_dispatch", boom)
    code = st.main(["palette", "x.png"])
    assert code == 1
    assert "測試用錯誤" in capsys.readouterr().err


def _write(tmp_path, name, text):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


def _two_tone_png(tmp_path):
    im = Image.new("RGB", (100, 50), "#EAF1FB")
    ImageDraw.Draw(im).rectangle([0, 0, 49, 49], fill="#2B579A")
    p = tmp_path / "src.png"
    im.save(p)
    return p


def test_palette_reports_size_and_dominant_colors(tmp_path):
    result = st.palette(_two_tone_png(tmp_path), n=8)
    assert result["size"] == (100, 50)
    hexes = [c["hex"] for c in result["colors"]]
    assert "#2B579A" in hexes
    assert "#EAF1FB" in hexes
    assert sum(c["ratio"] for c in result["colors"]) == pytest.approx(1.0, abs=0.001)


def test_palette_samples_exact_pixels(tmp_path):
    result = st.palette(_two_tone_png(tmp_path), at=["10,10", "80,10"])
    assert [s["hex"] for s in result["samples"]] == ["#2B579A", "#EAF1FB"]


def test_palette_rejects_out_of_range_point(tmp_path):
    with pytest.raises(st.SvgTraceError, match="超出圖片範圍"):
        st.palette(_two_tone_png(tmp_path), at=["999,999"])


def test_palette_missing_file_raises(tmp_path):
    with pytest.raises(st.SvgTraceError, match="圖片不存在"):
        st.palette(tmp_path / "nope.png")


def test_palette_clamps_n_to_max_256(tmp_path):
    result = st.palette(_two_tone_png(tmp_path), n=300)
    assert len(result["colors"]) <= 256
    assert "size" in result
    assert "colors" in result
    assert "samples" in result


def test_palette_rejects_malformed_coordinate(tmp_path):
    with pytest.raises(st.SvgTraceError, match="座標格式要寫成 X,Y"):
        st.palette(_two_tone_png(tmp_path), at=["abc"])
    with pytest.raises(st.SvgTraceError, match="座標格式要寫成 X,Y"):
        st.palette(_two_tone_png(tmp_path), at=["1,2,3"])


def test_svg_size_from_viewbox(tmp_path):
    p = _write(tmp_path, "a.svg", '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1184 666"/>')
    assert st.svg_size_px(p) == (1184.0, 666.0)


def test_svg_size_falls_back_to_width_height(tmp_path):
    p = _write(tmp_path, "b.svg", '<svg xmlns="http://www.w3.org/2000/svg" width="640" height="320"/>')
    assert st.svg_size_px(p) == (640.0, 320.0)


def test_svg_size_strips_units(tmp_path):
    p = _write(tmp_path, "c.svg", '<svg xmlns="http://www.w3.org/2000/svg" width="640px" height="320px"/>')
    assert st.svg_size_px(p) == (640.0, 320.0)


def test_svg_size_without_any_size_raises(tmp_path):
    p = _write(tmp_path, "d.svg", '<svg xmlns="http://www.w3.org/2000/svg"/>')
    with pytest.raises(st.SvgTraceError, match="viewBox"):
        st.svg_size_px(p)


def test_svg_size_parse_error_raises(tmp_path):
    p = _write(tmp_path, "e.svg", "<svg><unclosed>")
    with pytest.raises(st.SvgTraceError, match="解析失敗"):
        st.svg_size_px(p)


def test_svg_size_missing_file_raises(tmp_path):
    with pytest.raises(st.SvgTraceError, match="檔案不存在"):
        st.svg_size_px(tmp_path / "nope.svg")


def test_fit_box_height_limited():
    # 1184x666 塞進 12.333 x 6.5 吋: 高度先撞牆, 高吃滿 6.5, 寬按比例縮並水平置中
    x, y, w, h = st.fit_box(1184, 666, (0.5, 0.5, 12.333, 6.5))
    assert h == pytest.approx(6.5)
    assert w == pytest.approx(11.5556, abs=0.001)
    assert x == pytest.approx(0.8887, abs=0.001)
    assert y == pytest.approx(0.5)


def test_fit_box_width_limited():
    # 1184x400 塞進同一個 box: 寬度先撞牆, 寬吃滿 12.333, 高按比例縮並垂直置中
    x, y, w, h = st.fit_box(1184, 400, (0.5, 0.5, 12.333, 6.5))
    assert w == pytest.approx(12.333, abs=0.001)
    assert h == pytest.approx(4.1667, abs=0.001)
    assert x == pytest.approx(0.5)
    assert y == pytest.approx(1.6667, abs=0.001)


def test_svg_size_viewbox_token_count_wrong(tmp_path):
    p = _write(tmp_path, "f.svg", '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100"/>')
    with pytest.raises(st.SvgTraceError, match="viewBox"):
        st.svg_size_px(p)


def test_svg_size_viewbox_contains_non_numeric(tmp_path):
    p = _write(tmp_path, "g.svg", '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 abc 666"/>')
    with pytest.raises(st.SvgTraceError, match="viewBox"):
        st.svg_size_px(p)


def test_svg_size_invalid_width_value(tmp_path):
    p = _write(tmp_path, "h.svg", '<svg xmlns="http://www.w3.org/2000/svg" width="abc" height="10"/>')
    with pytest.raises(st.SvgTraceError, match="width.*無效"):
        st.svg_size_px(p)


def test_svg_size_invalid_height_value(tmp_path):
    p = _write(tmp_path, "i.svg", '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="xyz"/>')
    with pytest.raises(st.SvgTraceError, match="height.*無效"):
        st.svg_size_px(p)


SVG_OK = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200">
  <rect x="10" y="10" width="100" height="50" fill="#2B579A"/>
  <text x="20" y="40" font-size="16" fill="#FFFFFF">框</text>
</svg>'''


def test_svg_problems_empty_for_supported_svg(tmp_path):
    p = _write(tmp_path, "ok.svg", SVG_OK)
    assert st.svg_problems(p) == []


def test_svg_problems_reports_unsupported_element(tmp_path):
    p = _write(tmp_path, "bad.svg",
               '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10">'
               '<foreignObject width="1" height="1"/></svg>')
    problems = st.svg_problems(p)
    assert len(problems) == 1
    assert "foreignObject" in problems[0]


def test_svg_problems_reports_parse_error(tmp_path):
    p = _write(tmp_path, "broken.svg", "<svg><unclosed>")
    assert any("解析失敗" in m for m in st.svg_problems(p))


def test_svg_problems_missing_file(tmp_path):
    p = tmp_path / "nonexistent.svg"
    problems = st.svg_problems(p)
    assert len(problems) == 1
    assert "檔案不存在" in problems[0]


def test_build_creates_16_9_deck_when_missing(tmp_path):
    svg = _write(tmp_path, "d.svg", SVG_OK)
    out = tmp_path / "out.pptx"
    st.build(svg, out)
    prs = Presentation(str(out))
    assert prs.slide_width == Inches(13.333)
    assert prs.slide_height == Inches(7.5)
    assert len(prs.slides._sldIdLst) == 1


def test_build_appends_a_slide_to_existing_deck(tmp_path):
    svg = _write(tmp_path, "d.svg", SVG_OK)
    out = tmp_path / "out.pptx"
    st.build(svg, out, name="第一張")
    st.build(svg, out, name="第二張")
    slides = list(Presentation(str(out)).slides)
    assert len(slides) == 2
    # 原有那頁不可被動到
    assert [s.shapes[0].name for s in slides] == ["第一張", "第二張"]


def test_build_keeps_existing_slide_size(tmp_path):
    out = tmp_path / "four-three.pptx"
    Presentation().save(str(out))  # python-pptx 預設 10 x 7.5 吋
    svg = _write(tmp_path, "d.svg", SVG_OK)
    st.build(svg, out)
    after = Presentation(str(out))
    assert after.slide_width == Inches(10)
    assert after.slide_height == Inches(7.5)


def test_build_slide_has_only_the_diagram_group(tmp_path):
    svg = _write(tmp_path, "d.svg", SVG_OK)
    out = tmp_path / "out.pptx"
    st.build(svg, out, name="流程")
    slide = list(Presentation(str(out)).slides)[0]
    assert len(slide.shapes) == 1
    assert slide.shapes[0].name == "流程"


def test_build_shape_ids_do_not_collide(tmp_path):
    svg = _write(tmp_path, "d.svg", SVG_OK)
    out = tmp_path / "out.pptx"
    st.build(svg, out)
    st.build(svg, out)
    for slide in Presentation(str(out)).slides:
        ids = re.findall(r'<p:cNvPr id="(\d+)"', slide.shapes._spTree.xml)
        assert len(ids) == len(set(ids)), "同一頁裡的圖案編號重複了"


def test_build_refuses_unsupported_svg(tmp_path):
    svg = _write(tmp_path, "bad2.svg",
                 '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10">'
                 '<foreignObject width="1" height="1"/></svg>')
    with pytest.raises(st.SvgTraceError, match="foreignObject"):
        st.build(svg, tmp_path / "x.pptx")
    assert not (tmp_path / "x.pptx").exists(), "驗證沒過就不該產出半成品"


def test_build_reports_locked_file_in_chinese(tmp_path, monkeypatch):
    svg = _write(tmp_path, "d.svg", SVG_OK)

    def deny(self, path):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr("pptx.presentation.Presentation.save", deny)
    with pytest.raises(st.SvgTraceError, match="PowerPoint"):
        st.build(svg, tmp_path / "locked.pptx")


def test_build_group_geometry_matches_fit_box(tmp_path):
    # 4:3 簡報 (10 x 7.5 吋), 用來確認群組真的是用 fit_box 的結果擺放,
    # 不是誤用 gx/gy 對調或乾脆寫死 16:9 的 box
    out = tmp_path / "four-three.pptx"
    Presentation().save(str(out))
    svg = _write(tmp_path, "d.svg", SVG_OK)
    st.build(svg, out, name="流程")
    prs = Presentation(str(out))
    shape = list(prs.slides)[0].shapes[0]

    box = (
        st.MARGIN_IN,
        st.MARGIN_IN,
        prs.slide_width.inches - 2 * st.MARGIN_IN,
        prs.slide_height.inches - 2 * st.MARGIN_IN,
    )
    ex, ey, ew, eh = st.fit_box(*st.svg_size_px(svg), box)
    assert shape.left == pytest.approx(int(Inches(ex)), abs=2)
    assert shape.top == pytest.approx(int(Inches(ey)), abs=2)
    assert shape.width == pytest.approx(int(Inches(ew)), abs=2)
    assert shape.height == pytest.approx(int(Inches(eh)), abs=2)


NESTED_SVG = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200">
  <g id="alpha">
    <rect x="10" y="10" width="80" height="40" fill="#2B579A"/>
    <text x="20" y="35" font-size="14" fill="#FFFFFF">Alpha</text>
  </g>
  <g id="beta">
    <rect x="200" y="10" width="80" height="40" fill="#A52B2B"/>
    <text x="210" y="35" font-size="14" fill="#FFFFFF">Beta</text>
  </g>
</svg>'''


def test_build_preserves_nested_groups_and_text(tmp_path):
    svg = _write(tmp_path, "nested.svg", NESTED_SVG)
    out = tmp_path / "out.pptx"
    st.build(svg, out, name="流程圖")
    top = list(Presentation(str(out)).slides)[0].shapes[0]
    assert top.shape_type == MSO_SHAPE_TYPE.GROUP

    subgroups = list(top.shapes)
    assert len(subgroups) == 2
    assert all(sg.shape_type == MSO_SHAPE_TYPE.GROUP for sg in subgroups)

    texts = []
    for sg in subgroups:
        found = [sh.text_frame.text for sh in sg.shapes if sh.has_text_frame and sh.text_frame.text]
        assert len(found) == 1
        texts.append(found[0])
    assert texts == ["Alpha", "Beta"]


def test_build_uses_svg_group_id_as_shape_name(tmp_path):
    svg = _write(tmp_path, "named.svg", NESTED_SVG)
    out = tmp_path / "out.pptx"
    st.build(svg, out, name="流程圖")
    top = list(Presentation(str(out)).slides)[0].shapes[0]
    subgroups = list(top.shapes)
    assert [sg.name for sg in subgroups] == ["alpha", "beta"]


def test_find_chrome_prefers_env_var(tmp_path, monkeypatch):
    fake = tmp_path / "chrome.exe"
    fake.write_bytes(b"")
    monkeypatch.setenv("CHROME", str(fake))
    assert st.find_chrome() == fake


def test_find_chrome_lists_tried_paths_when_missing(tmp_path, monkeypatch):
    monkeypatch.delenv("CHROME", raising=False)
    monkeypatch.setattr(st, "CHROME_CANDIDATES", (str(tmp_path / "nope.exe"),))
    with pytest.raises(st.SvgTraceError, match="找不到 Chrome"):
        st.find_chrome()


def test_compare_puts_two_images_side_by_side(tmp_path):
    a, b = tmp_path / "a.png", tmp_path / "b.png"
    Image.new("RGB", (100, 50), "#FF0000").save(a)
    Image.new("RGB", (100, 50), "#00FF00").save(b)
    im = Image.open(st.compare(a, b, tmp_path / "cmp.png"))
    assert im.width == 100 + st.GUTTER_PX + 100
    assert im.height == 50 + st.LABEL_PX


def test_compare_scales_to_equal_height(tmp_path):
    a, b = tmp_path / "a.png", tmp_path / "b.png"
    Image.new("RGB", (100, 50), "#FF0000").save(a)
    Image.new("RGB", (200, 200), "#00FF00").save(b)
    im = Image.open(st.compare(a, b, tmp_path / "cmp2.png"))
    # 兩張都拉到 200 高: 原圖變 400 寬, 臨摹圖維持 200 寬
    assert im.width == 400 + st.GUTTER_PX + 200
    assert im.height == 200 + st.LABEL_PX


@pytest.mark.skipif(not st.chrome_available(), reason="這台機器沒有 Chrome 或 Edge")
def test_render_produces_a_png(tmp_path):
    svg = _write(tmp_path, "r.svg", SVG_OK)
    out = st.render(svg)
    assert out.exists() and out.stat().st_size > 0
    assert Image.open(out).size == (400, 200)


@pytest.mark.skipif(not st.chrome_available(), reason="這台機器沒有 Chrome 或 Edge")
def test_render_accepts_relative_output_path(tmp_path, monkeypatch):
    svg = _write(tmp_path, "r.svg", SVG_OK)
    monkeypatch.chdir(tmp_path)
    out = st.render(Path("r.svg"), Path("shot.png"))
    assert out.exists() and out.stat().st_size > 0


@pytest.mark.skipif(not st.chrome_available(), reason="這台機器沒有 Chrome 或 Edge")
def test_render_creates_missing_output_directory(tmp_path):
    svg = _write(tmp_path, "r.svg", SVG_OK)
    out_path = tmp_path / "sub" / "dir" / "shot.png"
    out = st.render(svg, out_path)
    assert out.exists() and out.stat().st_size > 0


def test_render_reports_svgtrace_error_not_attribute_error(tmp_path, monkeypatch):
    svg = _write(tmp_path, "r.svg", SVG_OK)
    monkeypatch.setattr(st, "find_chrome", lambda: Path("fake-chrome.exe"))

    class FakeProc:
        returncode = 1
        stderr = None

    monkeypatch.setattr(st.subprocess, "run", lambda *a, **k: FakeProc())
    with pytest.raises(st.SvgTraceError):
        st.render(svg)


def test_render_reports_timeout_as_svgtrace_error(tmp_path, monkeypatch):
    svg = _write(tmp_path, "r.svg", SVG_OK)
    monkeypatch.setattr(st, "find_chrome", lambda: Path("fake-chrome.exe"))

    def boom(*a, **k):
        raise st.subprocess.TimeoutExpired(cmd="fake-chrome.exe", timeout=st.RENDER_TIMEOUT_SEC)

    monkeypatch.setattr(st.subprocess, "run", boom)
    with pytest.raises(st.SvgTraceError, match="逾時"):
        st.render(svg)


class _FakeProc:
    def __init__(self, returncode, stderr=""):
        self.returncode = returncode
        self.stdout = ""
        self.stderr = stderr


def test_grab_reports_empty_clipboard(tmp_path, monkeypatch):
    monkeypatch.setattr(st.subprocess, "run", lambda *a, **k: _FakeProc(3))
    with pytest.raises(st.SvgTraceError, match="剪貼簿沒有圖片"):
        st.grab(tmp_path / "out.png")


def test_grab_reports_other_failures(tmp_path, monkeypatch):
    monkeypatch.setattr(st.subprocess, "run", lambda *a, **k: _FakeProc(1, "存取被拒"))
    with pytest.raises(st.SvgTraceError, match="取剪貼簿影像失敗"):
        st.grab(tmp_path / "out.png")


def test_grab_returns_path_on_success(tmp_path, monkeypatch):
    target = tmp_path / "sub" / "out.png"

    def fake_run(*a, **k):
        target.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (10, 10), "#FFFFFF").save(target)
        return _FakeProc(0)

    monkeypatch.setattr(st.subprocess, "run", fake_run)
    assert st.grab(target) == target


def test_work_dir_is_timestamped_under_home(tmp_path, monkeypatch):
    monkeypatch.setattr(st.Path, "home", staticmethod(lambda: tmp_path))
    d = st.work_dir()
    assert d.exists()
    assert d.parent == tmp_path / "svg-trace"
    assert re.fullmatch(r"\d{8}-\d{6}", d.name)


def test_grab_reports_svgtrace_error_not_attribute_error(tmp_path, monkeypatch):
    class FakeProc:
        returncode = 1
        stderr = None

    monkeypatch.setattr(st.subprocess, "run", lambda *a, **k: FakeProc())
    with pytest.raises(st.SvgTraceError):
        st.grab(tmp_path / "out.png")


def test_grab_reports_timeout_as_svgtrace_error(tmp_path, monkeypatch):
    def boom(*a, **k):
        raise st.subprocess.TimeoutExpired(cmd="powershell.exe", timeout=st.GRAB_TIMEOUT_SEC)

    monkeypatch.setattr(st.subprocess, "run", boom)
    with pytest.raises(st.SvgTraceError, match="逾時"):
        st.grab(tmp_path / "out.png")


def test_svg_problems_reports_count_when_exceeds_limit(tmp_path):
    # 造 9 個 foreignObject, 超過 8 個上限
    svg_with_many = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">
      <foreignObject x="0" y="0" width="10" height="10"/>
      <foreignObject x="10" y="0" width="10" height="10"/>
      <foreignObject x="20" y="0" width="10" height="10"/>
      <foreignObject x="30" y="0" width="10" height="10"/>
      <foreignObject x="40" y="0" width="10" height="10"/>
      <foreignObject x="50" y="0" width="10" height="10"/>
      <foreignObject x="60" y="0" width="10" height="10"/>
      <foreignObject x="70" y="0" width="10" height="10"/>
      <foreignObject x="80" y="0" width="10" height="10"/>
    </svg>'''
    p = _write(tmp_path, "many.svg", svg_with_many)
    problems = st.svg_problems(p)
    assert len(problems) == 1
    assert "foreignObject" in problems[0]
    assert "還有 1 個" in problems[0]
