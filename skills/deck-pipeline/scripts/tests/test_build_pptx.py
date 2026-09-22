import re
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import build_pptx as bp  # noqa: E402

SKILL = Path(__file__).resolve().parents[2]
MINI = SKILL / "examples" / "mini-deck"

SLIDES = """# 測試簡報

## 01 背景

### [title] 為什麼
副標一句

### [text] 三個痛點
- 痛點一
- 痛點二

> 講稿對應: 第一段
> 第二行

### [image] 舊流程
![舊流程截圖](01_背景/flow_old.png)
- 一句註解

## 02 做法

### [compare] 舊 vs 新
| 舊 | 新 |
|---|---|
| 手動 | 自動 |

### [image-text] 理想
![描述: 箭頭圖](TODO)
- 六階段
"""


def test_parse_slides_title_and_topics():
    deck = bp.parse_slides(SLIDES)
    assert deck.title == "測試簡報"
    assert [t.folder for t in deck.topics] == ["01_背景", "02_做法"]
    assert [len(t.pages) for t in deck.topics] == [3, 2]


def test_parse_slides_page_fields():
    deck = bp.parse_slides(SLIDES)
    p1, p2, p3 = deck.topics[0].pages
    assert (p1.layout, p1.title, p1.subtitle) == ("title", "為什麼", "副標一句")
    assert p2.bullets == ["痛點一", "痛點二"]
    assert p2.notes == "講稿對應: 第一段\n第二行"
    assert p3.images == [("舊流程截圖", "01_背景/flow_old.png")]
    assert p3.bullets == ["一句註解"]


def test_parse_slides_table_and_todo():
    deck = bp.parse_slides(SLIDES)
    cmp_, img = deck.topics[1].pages
    assert cmp_.table == [["舊", "新"], ["手動", "自動"]]
    assert img.images == [("描述: 箭頭圖", bp.TODO)]


def test_parse_slides_pages_iter():
    deck = bp.parse_slides(SLIDES)
    seq = [(t.num, i, p.layout) for t, i, p in deck.pages()]
    assert seq == [("01", 1, "title"), ("01", 2, "text"), ("01", 3, "image"),
                   ("02", 1, "compare"), ("02", 2, "image-text")]


def test_parse_slides_rejects_page_without_layout():
    with pytest.raises(bp.SlidesParseError, match="版型"):
        bp.parse_slides("# t\n\n## 01 a\n\n### 沒有標記\n")


def test_parse_slides_rejects_missing_title():
    with pytest.raises(bp.SlidesParseError, match="簡報名稱"):
        bp.parse_slides("## 01 a\n\n### [text] x\n")


STYLE = """# 版型

說明文字, 腳本忽略。

```yaml
slide: {w: 13.333, h: 7.5}
theme:
  bg: "#FFFFFF"
  fg: "#1F2937"
  accent: "#2563EB"
  font_title: "Noto Sans TC"
  font_body: "Noto Sans TC"
layouts:
  title:
    - {role: title, box: [1.0, 2.4, 11.3, 1.4], size: 40, bold: true}
    - {role: subtitle, box: [1.0, 4.0, 11.3, 1.0], size: 22}
  text:
    - {role: title, box: [0.8, 0.6, 11.7, 1.2], size: 32, bold: true}
    - {role: body, box: [0.8, 2.2, 11.7, 4.5], size: 20}
```
"""


def test_parse_style_reads_yaml_block():
    st = bp.parse_style(STYLE)
    assert st.slide == {"w": 13.333, "h": 7.5}
    assert st.theme["accent"] == "#2563EB"
    assert set(st.layouts) == {"title", "text"}
    assert st.layouts["text"][1]["role"] == "body"


def test_parse_style_rejects_missing_yaml():
    with pytest.raises(bp.StyleParseError, match="yaml"):
        bp.parse_style("# 只有文字沒有區塊\n")


def test_parse_style_rejects_missing_key():
    with pytest.raises(bp.StyleParseError, match="layouts"):
        bp.parse_style("```yaml\nslide: {w: 1, h: 1}\ntheme: {}\n```\n")


def test_validate_reports_unknown_layout_and_missing_image(tmp_path):
    deck = bp.parse_slides(SLIDES)   # 用到 image / compare / image-text, STYLE 只定義 title / text
    st = bp.parse_style(STYLE)
    errors = bp.validate(deck, st, tmp_path)
    assert any("01_背景 第3頁 [image]" in e and "版型未在 STYLE.md 定義" in e for e in errors)
    assert any("圖片不存在 01_背景/flow_old.png" in e for e in errors)
    assert not any("TODO" in e for e in errors)   # TODO 不算缺圖


def test_validate_passes_on_mini_deck():
    deck = bp.parse_slides((MINI / "SLIDES.md").read_text(encoding="utf-8"))
    st = bp.parse_style((MINI / "STYLE.md").read_text(encoding="utf-8"))
    assert bp.validate(deck, st, MINI) == []


def _build_mini(tmp_path):
    deck = bp.parse_slides((MINI / "SLIDES.md").read_text(encoding="utf-8"))
    st = bp.parse_style((MINI / "STYLE.md").read_text(encoding="utf-8"))
    prs = bp.render(deck, st, MINI)
    out = tmp_path / "mini.pptx"
    prs.save(str(out))
    return Presentation(str(out))


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_render_slide_count_and_size(tmp_path):
    prs = _build_mini(tmp_path)
    assert len(prs.slides) == 6
    assert prs.slide_width == Inches(13.333)
    assert prs.slide_height == Inches(7.5)


def test_render_title_and_bullets_are_editable_text(tmp_path):
    prs = _build_mini(tmp_path)
    s1, s2 = prs.slides[0], prs.slides[1]
    assert "為什麼要做這件事" in _texts(s1)
    assert "把散亂素材變成一份簡報" in _texts(s1)
    assert "素材散在各處\n講稿與頁面混在一起\n排版每次從零開始" in _texts(s2)


def test_render_notes(tmp_path):
    prs = _build_mini(tmp_path)
    assert prs.slides[1].notes_slide.notes_text_frame.text == "講稿對應: 背景第一段"
    assert not prs.slides[0].has_notes_slide


def test_render_picture_fits_box(tmp_path):
    prs = _build_mini(tmp_path)
    pics = [sh for sh in prs.slides[2].shapes if sh.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
    assert len(pics) == 1
    pic = pics[0]
    # box [0.8, 1.8, 11.7, 4.6], 圖 400x300 → 高度吃滿 4.6, 寬 6.133, 水平置中
    assert abs(pic.height - Inches(4.6)) < Inches(0.01)
    assert abs(pic.width - Inches(4.6 * 400 / 300)) < Inches(0.01)
    assert abs(pic.left - Inches(0.8 + (11.7 - 4.6 * 400 / 300) / 2)) < Inches(0.01)


def test_render_todo_image_becomes_gray_placeholder(tmp_path):
    prs = _build_mini(tmp_path)
    s5 = prs.slides[4]  # 02 第2頁 image-text
    rects = [sh for sh in s5.shapes if sh.shape_type == 1]  # AUTO_SHAPE
    assert len(rects) == 1
    assert rects[0].fill.fore_color.rgb == bp.RGBColor(0xD1, 0xD5, 0xDB)
    assert "[待補圖] 描述: 一張簡化的箭頭圖, 三步變一步" in rects[0].text_frame.text


def test_render_compare_columns(tmp_path):
    prs = _build_mini(tmp_path)
    texts = _texts(prs.slides[5])
    assert "舊\n手動搬檔\n邊寫邊排" in texts
    assert "新\nAI 分類\n先講稿再頁面" in texts


def test_safe_filename():
    assert bp.safe_filename('a/b:c*d?"e<f>g|h') == "a_b_c_d_e_f_g_h"
    assert bp.safe_filename("  ") == "deck"


def test_main_builds_to_output(tmp_path):
    import shutil
    shutil.copytree(MINI, tmp_path / "deck", ignore=shutil.ignore_patterns("output"))
    assert bp.main([str(tmp_path / "deck")]) == 0
    assert (tmp_path / "deck" / "output" / "mini-deck 範例.pptx").is_file()


def test_main_returns_1_on_validation_error(tmp_path, capsys):
    import shutil
    shutil.copytree(MINI, tmp_path / "deck", ignore=shutil.ignore_patterns("output"))
    (tmp_path / "deck" / "01_背景" / "flow_old.png").unlink()
    assert bp.main([str(tmp_path / "deck")]) == 1
    assert "圖片不存在" in capsys.readouterr().out
    assert not (tmp_path / "deck" / "output").exists()


def test_images_todo_lists_only_todo_images():
    deck = bp.parse_slides(SLIDES)
    md = bp.images_todo(deck)
    assert md.startswith("# 待補圖清單")
    assert "| 02_做法 第2頁 理想 | 描述: 箭頭圖 | 02_做法/p02_img1.png |" in md
    assert "flow_old" not in md


def test_images_todo_empty():
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [text] x\n- y\n")
    assert "| (無) | | |" in bp.images_todo(deck)


def test_main_images_todo_flag(tmp_path):
    import shutil
    shutil.copytree(MINI, tmp_path / "deck", ignore=shutil.ignore_patterns("output"))
    assert bp.main([str(tmp_path / "deck"), "--images-todo"]) == 0
    md = (tmp_path / "deck" / "IMAGES_TODO.md").read_text(encoding="utf-8")
    assert "02_做法/p02_img1.png" in md
    assert not (tmp_path / "deck" / "output").exists()


# F1: main 印乾淨錯誤訊息, 不印 traceback


def test_main_reports_missing_slides_file_without_traceback(tmp_path, capsys):
    deck_dir = tmp_path / "deck"
    deck_dir.mkdir()
    assert bp.main([str(deck_dir)]) == 1
    out = capsys.readouterr().out
    assert "build 失敗" in out
    assert "SLIDES.md" in out
    assert "Traceback" not in out


def test_main_reports_slides_parse_error_without_traceback(tmp_path, capsys):
    deck_dir = tmp_path / "deck"
    deck_dir.mkdir()
    (deck_dir / "SLIDES.md").write_text("## 01 a\n\n### [text] x\n", encoding="utf-8")
    assert bp.main([str(deck_dir)]) == 1
    out = capsys.readouterr().out
    assert "build 失敗" in out
    assert "簡報名稱" in out
    assert "Traceback" not in out


# F2: 格式錯誤的 ## 標題要報錯, 不能悄悄變成 subtitle


def test_parse_slides_rejects_topic_heading_without_two_digit_number():
    with pytest.raises(bp.SlidesParseError, match="topic 標題"):
        bp.parse_slides("# t\n\n## 背景\n\n### [text] x\n")


def test_parse_slides_rejects_topic_heading_with_single_digit_number():
    with pytest.raises(bp.SlidesParseError, match="topic 標題"):
        bp.parse_slides("# t\n\n## 1 a\n")


# F3: 文字框不自動縮放, 且尺寸照 STYLE.md box


def test_render_textboxes_do_not_autosize(tmp_path):
    prs = _build_mini(tmp_path)
    s1 = prs.slides[0]
    for sh in s1.shapes:
        if sh.has_text_frame:
            assert sh.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    title_shape = next(
        sh for sh in s1.shapes if sh.has_text_frame and sh.text_frame.text == "為什麼要做這件事"
    )
    assert abs(title_shape.height - Inches(1.4)) < Inches(0.01)


# F4: validate 先檢查 STYLE.md 結構


STYLE_INVALID = """```yaml
slide: {w: 13.333, h: 7.5}
theme:
  bg: "#FFFFFF"
  fg: "#1F2937"
  font_body: "Noto Sans TC"
layouts:
  bad:
    - {role: footer, box: [1.0, 2.4, 11.3, 1.4]}
    - {role: title}
```
"""


def test_validate_checks_style_structure(tmp_path):
    deck = bp.Deck(title="t")
    st = bp.parse_style(STYLE_INVALID)
    errors = bp.validate(deck, st, tmp_path)
    assert len(errors) == 3
    assert any("role 不合法: footer" in e for e in errors)
    assert any("box 必須是 4 個數字" in e for e in errors)
    assert any("theme 缺少 font_title" in e for e in errors)


def test_validate_style_structure_still_passes_on_mini_deck():
    deck = bp.parse_slides((MINI / "SLIDES.md").read_text(encoding="utf-8"))
    st = bp.parse_style((MINI / "STYLE.md").read_text(encoding="utf-8"))
    assert bp.validate(deck, st, MINI) == []


# F6: > 後面沒空格也算 note


def test_parse_slides_notes_without_space_after_gt():
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [text] x\n>講稿對應: 第一段\n")
    page = deck.topics[0].pages[0]
    assert page.notes == "講稿對應: 第一段"
    assert page.subtitle == ""


# F7: RE_YAML 錨定行首, CRLF 仍可 parse


def test_parse_style_handles_crlf():
    text = (MINI / "STYLE.md").read_text(encoding="utf-8")
    crlf_text = text.replace("\n", "\r\n")
    st = bp.parse_style(crlf_text)
    assert set(st.layouts) == {"title", "section", "text", "image", "image-text", "compare"}


TABLE_SLIDES = """# t

## 01 a

### [table] 開發時程對照
| 項目 | 行數 | 天數 |
| 分組日榜 | 14,140 | 15 |
| 通行證 | 16,904 | 36 |
- 規模接近, 時間少一半
"""

TABLE_STYLE = """```yaml
slide: {w: 13.333, h: 7.5}
theme: {bg: "#FFFFFF", fg: "#1F2937", accent: "#2563EB", font_title: "Noto Sans TC", font_body: "Noto Sans TC"}
layouts:
  table:
    - {role: title, box: [0.8, 0.5, 11.7, 1.0], size: 28, bold: true}
    - {role: table, box: [0.8, 1.7, 11.7, 3.0], size: 18}
    - {role: body,  box: [0.8, 5.2, 11.7, 0.8], size: 20}
```
"""


def test_validate_accepts_table_role(tmp_path):
    deck = bp.parse_slides(TABLE_SLIDES)
    st = bp.parse_style(TABLE_STYLE)
    assert bp.validate(deck, st, tmp_path) == []


def test_render_table_role_builds_real_table(tmp_path):
    deck = bp.parse_slides(TABLE_SLIDES)
    st = bp.parse_style(TABLE_STYLE)
    prs = bp.render(deck, st, tmp_path)
    out = tmp_path / "t.pptx"
    prs.save(str(out))
    slide = Presentation(str(out)).slides[0]
    frames = [sh for sh in slide.shapes if sh.has_table]
    assert len(frames) == 1
    tbl = frames[0].table
    assert len(tbl.rows) == 3 and len(tbl.columns) == 3
    assert tbl.cell(0, 0).text == "項目"
    assert tbl.cell(2, 1).text == "16,904"
    header_run = tbl.cell(0, 0).text_frame.paragraphs[0].runs[0]
    body_run = tbl.cell(1, 0).text_frame.paragraphs[0].runs[0]
    assert header_run.font.bold is True
    assert body_run.font.bold is False
    assert tbl.cell(0, 0).fill.fore_color.rgb == bp.RGBColor(0x25, 0x63, 0xEB)
    assert tbl.cell(1, 0).fill.fore_color.rgb == bp.RGBColor(0xFF, 0xFF, 0xFF)
    assert "規模接近, 時間少一半" in [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_parse_slides_image_path_with_parentheses():
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [image] x\n![desc](01_a(skill-name)/shot (1).png)\n")
    assert deck.topics[0].pages[0].images == [("desc", "01_a(skill-name)/shot (1).png")]


SVG_MINI = """<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200" viewBox="0 0 400 200">
  <defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="9" refY="5" orient="auto"><path d="M0,0 L10,5 L0,10 z" fill="#B45309"/></marker></defs>
  <rect x="20" y="60" width="120" height="60" rx="8" fill="#EFE8DE" stroke="#2B2622"/>
  <text x="80" y="95" font-size="16" text-anchor="middle" fill="#2B2622">A</text>
  <line x1="140" y1="90" x2="260" y2="90" stroke="#B45309" stroke-dasharray="6 4" marker-end="url(#arrow)"/>
  <rect x="260" y="60" width="120" height="60" rx="8" fill="#EFE8DE" stroke="#2B2622"/>
</svg>
"""


def _svg_deck(tmp_path):
    (tmp_path / "01_a").mkdir()
    (tmp_path / "01_a" / "d.svg").write_text(SVG_MINI, encoding="utf-8")
    slides = "# t\n\n## 01 a\n\n### [image] x\n![圖案圖](01_a/d.svg)\n- 一句\n"
    style = """```yaml
slide: {w: 13.333, h: 7.5}
theme: {bg: "#FFFFFF", fg: "#1F2937", accent: "#2563EB", font_title: "Noto Sans TC", font_body: "Noto Sans TC"}
layouts:
  image:
    - {role: title, box: [0.8, 0.5, 11.7, 1.0], size: 28, bold: true}
    - {role: image, box: [0.8, 1.8, 11.7, 4.6]}
    - {role: body,  box: [0.8, 6.6, 11.7, 0.6], size: 16}
```
"""
    return bp.parse_slides(slides), bp.parse_style(style)


def test_validate_accepts_svg_image(tmp_path):
    deck, st = _svg_deck(tmp_path)
    assert bp.validate(deck, st, tmp_path) == []


def test_render_svg_becomes_native_shape_group_fitted_in_box(tmp_path):
    deck, st = _svg_deck(tmp_path)
    prs = bp.render(deck, st, tmp_path)
    out = tmp_path / "t.pptx"
    prs.save(str(out))
    slide = Presentation(str(out)).slides[0]
    groups = [sh for sh in slide.shapes if sh.shape_type == 6]  # MSO_SHAPE_TYPE.GROUP
    assert len(groups) == 1
    grp = groups[0]
    assert len(grp.shapes) == 4  # rect, text, line, rect
    # box [0.8, 1.8, 11.7, 4.6], svg 400x200 (2:1) → 寬度受限 11.7 → 高 5.85 > 4.6, 所以高度吃滿 4.6, 寬 9.2, 水平置中
    assert abs(grp.height - Inches(4.6)) < Inches(0.01)
    assert abs(grp.width - Inches(9.2)) < Inches(0.01)
    assert abs(grp.left - Inches(0.8 + (11.7 - 9.2) / 2)) < Inches(0.01)
    assert abs(grp.top - Inches(1.8)) < Inches(0.01)
    assert grp.name == "diagram d"


def test_render_svg_shape_ids_do_not_collide(tmp_path):
    import zipfile
    deck, st = _svg_deck(tmp_path)
    prs = bp.render(deck, st, tmp_path)
    out = tmp_path / "t.pptx"
    prs.save(str(out))
    with zipfile.ZipFile(out) as z:
        xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
    ids = re.findall(r'<p:cNvPr id="(\d+)"', xml)
    assert len(ids) == len(set(ids)), ids


def test_validate_reports_svg_with_unsupported_element(tmp_path):
    deck, st = _svg_deck(tmp_path)
    (tmp_path / "01_a" / "d.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><foreignObject width="1" height="1"/></svg>', encoding="utf-8")
    errors = bp.validate(deck, st, tmp_path)
    assert any("SVG" in e and "foreignObject" in e for e in errors)


# F9: <!-- --> 註解行不進投影片, 給人看的頁碼標記用


def test_parse_slides_ignores_html_comments():
    deck = bp.parse_slides(
        "# t\n\n## 01 a\n\n<!-- P01 -->\n### [text] x\n<!-- 分隔 -->\n- 條列\n"
    )
    page = deck.topics[0].pages[0]
    assert page.subtitle == ""
    assert page.bullets == ["條列"]
    assert len(deck.topics[0].pages) == 1


# F10: ### 只放版型時, 下一行純文字是標題; 標題可帶 [章節標籤] 前綴


def test_parse_slides_title_on_own_line_with_label():
    deck = bp.parse_slides(
        "# t\n\n## 01 a\n\n### [image]\n[小結] 團隊擴大最明顯的變化\n管理職不只是用來管人\n![描述](TODO)\n"
    )
    page = deck.topics[0].pages[0]
    assert page.label == "小結"
    assert page.title == "團隊擴大最明顯的變化"
    assert page.subtitle == "管理職不只是用來管人"
    assert page.images == [("描述", bp.TODO)]


def test_parse_slides_title_on_own_line_without_label():
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [text]\n純標題\n副標\n")
    page = deck.topics[0].pages[0]
    assert (page.label, page.title, page.subtitle) == ("", "純標題", "副標")


def test_parse_slides_old_inline_title_still_works():
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [text] 舊寫法標題\n副標\n")
    page = deck.topics[0].pages[0]
    assert (page.label, page.title, page.subtitle) == ("", "舊寫法標題", "副標")


def test_validate_accepts_label_role():
    assert "label" in bp.VALID_ROLES


# F11: cards role — page.table 每列一張卡 (| 徽章 | 標題 | 說明 |)


def test_cards_role_is_valid():
    assert "cards" in bp.VALID_ROLES


def test_render_cards_layout(tmp_path):
    slides = (
        "# t\n\n## 01 a\n\n### [cards]\n能做到什麼\n"
        "| 現成可用 | 回合流程 | 發牌抽牌比大小 |\n"
        "| 要自己寫 | 下注輪與底池 | 盲注邊池最小加注都沒有 |\n"
        "| 沒有 | 節點編輯器 | 文件與截圖都找不到畫布 |\n"
    )
    style = (
        "```yaml\n"
        "slide: {w: 13.333, h: 7.5}\n"
        "theme: {bg: '#FFFFFF', fg: '#1F2937', accent: '#2563EB', font_title: A, font_body: B}\n"
        "layouts:\n"
        "  cards:\n"
        "    - {role: title, box: [0.8, 0.5, 11.7, 1.0], size: 28, bold: true}\n"
        "    - {role: cards, box: [0.8, 1.8, 11.7, 4.6], size: 16, cols: 3,\n"
        "       badges: {現成可用: '#1F2937', 要自己寫: '#9CA3AF', 沒有: '#DC2626'}}\n"
        "```\n"
    )
    deck = bp.parse_slides(slides)
    st = bp.parse_style(style)
    assert bp.validate(deck, st, tmp_path) == []
    prs = bp.render(deck, st, tmp_path)
    shapes = list(prs.slides[0].shapes)
    texts = [s.text_frame.text for s in shapes if s.has_text_frame]
    assert "01" in texts and "03" in texts
    assert "現成可用" in texts and "沒有" in texts
    assert "下注輪與底池" in texts


# F12: 帶 EXIF orientation 的照片 PowerPoint 不會轉正, validate 要擋下來


def _write_jpeg(path, size, orientation=None):
    from PIL import Image
    im = Image.new("RGB", size, (120, 140, 160))
    kwargs = {}
    if orientation is not None:
        ex = im.getexif()
        ex[274] = orientation
        kwargs["exif"] = ex.tobytes()
    im.save(path, "JPEG", **kwargs)


def test_exif_problem_flags_rotated_photo(tmp_path):
    f = tmp_path / "rotated.jpg"
    _write_jpeg(f, (400, 300), orientation=6)
    msg = bp._exif_problem(f)
    assert msg is not None and "orientation=6" in msg


def test_exif_problem_silent_on_upright_photo(tmp_path):
    upright = tmp_path / "upright.jpg"
    _write_jpeg(upright, (400, 300))
    tagged = tmp_path / "tagged.jpg"
    _write_jpeg(tagged, (400, 300), orientation=1)
    assert bp._exif_problem(upright) is None
    assert bp._exif_problem(tagged) is None


def test_validate_reports_rotated_photo(tmp_path):
    (tmp_path / "01_a").mkdir()
    _write_jpeg(tmp_path / "01_a" / "p.jpg", (400, 300), orientation=6)
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [image]\nx\n![d](01_a/p.jpg)\n")
    st = bp.parse_style(
        "```yaml\nslide: {w: 13.333, h: 7.5}\n"
        "theme: {bg: '#FFF', fg: '#000', font_title: A, font_body: B}\n"
        "layouts:\n  image:\n    - {role: image, box: [1, 1, 5, 4]}\n```\n"
    )
    errors = bp.validate(deck, st, tmp_path)
    assert any("EXIF orientation" in e for e in errors)


# F13: decor role — 版型固定的裝飾 (色塊 / 箭頭 / 疊在圖上的標籤), 不吃頁面內容


def test_decor_role_is_valid():
    assert "decor" in bp.VALID_ROLES


def test_render_decor_shapes(tmp_path):
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [section]\n章節標題\n")
    st = bp.parse_style(
        "```yaml\n"
        "slide: {w: 13.333, h: 7.5}\n"
        "theme: {bg: '#FFFFFF', fg: '#1F2937', accent: '#2563EB', font_title: A, font_body: B}\n"
        "layouts:\n"
        "  section:\n"
        "    - {role: decor, box: [0.75, 2.73, 0.79, 0.04]}\n"
        "    - {role: decor, box: [6.4, 3.0, 0.5, 0.4], shape: arrow, color: '#6B7280'}\n"
        "    - {role: decor, box: [1.0, 4.0, 1.0, 0.4], shape: text, text: 產出, size: 14}\n"
        "    - {role: title, box: [0.75, 3.13, 11.8, 1.9], size: 45, bold: true}\n"
        "```\n"
    )
    assert bp.validate(deck, st, tmp_path) == []
    prs = bp.render(deck, st, tmp_path)
    shapes = list(prs.slides[0].shapes)
    assert len(shapes) == 4
    texts = [s.text_frame.text for s in shapes if s.has_text_frame]
    assert "產出" in texts and "章節標題" in texts


# F14: 圖片描述開頭的 @photo / @ai / @svg 來源標籤, --images-todo 依此分組


def test_split_src_tag():
    assert bp.split_src_tag("@ai 意象: 沙堡") == ("ai", "意象: 沙堡")
    assert bp.split_src_tag("@photo 現場照片") == ("photo", "現場照片")
    assert bp.split_src_tag("@svg 流程圖") == ("svg", "流程圖")
    assert bp.split_src_tag("沒有標籤的描述") == ("", "沒有標籤的描述")
    assert bp.split_src_tag("@unknown 不是合法標籤") == ("", "@unknown 不是合法標籤")


def test_images_todo_groups_by_src_tag():
    deck = bp.parse_slides(
        "# t\n\n## 01 a\n\n### [image]\n第一頁\n![@ai 沙堡](TODO)\n"
        "\n### [image]\n第二頁\n![@svg 流程圖](TODO)\n"
        "\n### [image]\n第三頁\n![@photo 現場照](TODO)\n"
        "\n### [image]\n第四頁\n![沒標](TODO)\n"
    )
    md = bp.images_todo(deck)
    for marker in ("## @photo", "## @ai", "## @svg", "## (未標籤)"):
        assert marker in md, marker
    assert "沙堡" in md and "@ai 沙堡" not in md      # 標籤已從描述剝掉
    assert "回 SLIDES.md 在描述開頭補上" in md


def test_images_todo_empty():
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [image]\nx\n![d](01_a/p.png)\n")
    assert "目前沒有待補圖" in bp.images_todo(deck)


# F15: 圖片嵌進 pptx 前依顯示尺寸縮到目標 dpi, 並選較小的編碼


def _png(path, size, mode="RGB"):
    from PIL import Image
    import random
    im = Image.new(mode, size)
    px = im.load()
    random.seed(1)
    for y in range(0, size[1], 4):          # 隨機雜訊, 免得 PNG 壓到極小失去比較意義
        for x in range(0, size[0], 4):
            v = random.randrange(256)
            for dy in range(4):
                for dx in range(4):
                    if x+dx < size[0] and y+dy < size[1]:
                        px[x+dx, y+dy] = (v, v, v) if mode == "RGB" else (v, v, v, 255)
    im.save(path, "PNG")


def test_optimized_blob_downscales_to_dpi(tmp_path):
    f = tmp_path / "big.png"
    _png(f, (1920, 1080))
    buf = bp._optimized_blob(f, disp_w_in=5.71, export={"image_dpi": 150, "jpeg_quality": 88})
    from PIL import Image
    assert Image.open(buf).size[0] == round(5.71 * 150)


def test_optimized_blob_never_upscales(tmp_path):
    f = tmp_path / "small.png"
    _png(f, (200, 100))
    buf = bp._optimized_blob(f, disp_w_in=8.0, export=bp.DEFAULT_EXPORT)
    from PIL import Image
    assert buf is None or Image.open(buf).size[0] == 200


def _flat_png(path, size, mode="RGB"):
    """大尺寸但內容簡單的圖 — 縮放後一定壓得贏原檔"""
    from PIL import Image, ImageDraw
    im = Image.new(mode, size, (250, 250, 250) if mode == "RGB" else (250, 250, 250, 255))
    d = ImageDraw.Draw(im)
    for k in range(6):
        d.rectangle([k * 90, k * 40, k * 90 + 300, k * 40 + 200],
                    fill=(40 * k, 90, 200 - 20 * k) if mode == "RGB" else (40 * k, 90, 200 - 20 * k, 255))
    im.save(path, "PNG")


def test_optimized_blob_keeps_png_when_alpha(tmp_path):
    f = tmp_path / "alpha.png"
    _flat_png(f, (1600, 900), mode="RGBA")
    buf = bp._optimized_blob(f, disp_w_in=4.0, export=bp.DEFAULT_EXPORT)
    from PIL import Image
    assert buf is not None
    assert Image.open(buf).format == "PNG"      # 有 alpha 不可以轉成 JPEG


def test_optimized_blob_disabled_by_zero_dpi(tmp_path):
    f = tmp_path / "x.png"
    _png(f, (1920, 1080))
    assert bp._optimized_blob(f, 5.0, {"image_dpi": 0}) is None


def test_parse_style_export_defaults_and_override():
    base = ("```yaml\nslide: {w: 13.333, h: 7.5}\n"
            "theme: {bg: '#FFF', fg: '#000', font_title: A, font_body: B}\n"
            "layouts: {image: [{role: image, box: [1,1,5,4]}]}\n")
    assert bp.parse_style(base + "```\n").export == bp.DEFAULT_EXPORT
    st = bp.parse_style(base + "export: {image_dpi: 220}\n```\n")
    assert st.export["image_dpi"] == 220 and st.export["jpeg_quality"] == bp.DEFAULT_EXPORT["jpeg_quality"]


def test_render_optimize_flag_shrinks_output(tmp_path):
    (tmp_path / "01_a").mkdir()
    _flat_png(tmp_path / "01_a" / "p.png", (1920, 1080))
    deck = bp.parse_slides("# t\n\n## 01 a\n\n### [image]\nx\n![d](01_a/p.png)\n")
    st = bp.parse_style(
        "```yaml\nslide: {w: 13.333, h: 7.5}\n"
        "theme: {bg: '#FFFFFF', fg: '#000000', font_title: A, font_body: B}\n"
        "layouts: {image: [{role: image, box: [0.75, 2.08, 5.71, 2.25]}]}\n```\n")
    import zipfile
    big, small = tmp_path / "big.pptx", tmp_path / "small.pptx"
    bp.render(deck, st, tmp_path, optimize=False).save(str(big))
    bp.render(deck, st, tmp_path, optimize=True).save(str(small))

    def media_bytes(f):
        with zipfile.ZipFile(f) as z:
            return sum(i.file_size for i in z.infolist() if "/media/" in i.filename)

    # 比 media 而不是 pptx 檔案大小 — 小圖進 zip 會再被壓一次, 蓋掉差異
    assert media_bytes(small) < media_bytes(big)
