"""deck-pipeline: SLIDES.md + STYLE.md -> 可編輯 pptx

用法:
  py -3.12-64 build_pptx.py <deck資料夾>               # 產 output/<title>.pptx
  py -3.12-64 build_pptx.py <deck資料夾> --images-todo # 只產 IMAGES_TODO.md
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent / "vendor"))
from svg_to_pptx import ConvertContext, EMU_PER_PX, collect_defs, convert_element  # noqa: E402
from svg_to_pptx.drawingml_converter import _collect_unsupported_visuals  # noqa: E402
from svg_to_pptx.drawingml_utils import SVG_NS  # noqa: E402

TODO = "TODO"
SVG_SUFFIX = ".svg"


class SlidesParseError(ValueError):
    pass


@dataclass
class Page:
    layout: str
    title: str
    label: str = ""          # 左上角章節標籤 (案例2 / 小結), 標題行 [xxx] 前綴
    subtitle: str = ""
    bullets: list[str] = field(default_factory=list)
    images: list[tuple[str, str]] = field(default_factory=list)  # (描述, 路徑)
    table: list[list[str]] = field(default_factory=list)          # 含表頭列
    notes: str = ""


@dataclass
class Topic:
    num: str
    name: str
    pages: list[Page] = field(default_factory=list)

    @property
    def folder(self) -> str:
        return f"{self.num}_{self.name}"


@dataclass
class Deck:
    title: str
    topics: list[Topic] = field(default_factory=list)

    def pages(self):
        for topic in self.topics:
            for i, page in enumerate(topic.pages, 1):
                yield topic, i, page


RE_TOPIC = re.compile(r"^## (\d{2})\s+(.+?)\s*$")
RE_PAGE = re.compile(r"^### \[([\w-]+)\]\s*(.*?)\s*$")
RE_LABEL = re.compile(r"^\[([^\]]+)\]\s*(.+?)\s*$")
RE_IMAGE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
RE_TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}")


def parse_slides(text: str) -> Deck:
    deck = Deck(title="")
    topic: Topic | None = None
    page: Page | None = None
    for lineno, raw in enumerate(text.splitlines(), 1):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("<!--"):
            continue  # 註解行 (頁碼、分隔線等給人看的標記), 不進投影片
        if line.startswith("# "):
            deck.title = line[2:].strip()
            continue
        m = RE_TOPIC.match(line)
        if m:
            topic = Topic(m.group(1), m.group(2))
            deck.topics.append(topic)
            page = None
            continue
        if line.startswith("## "):
            raise SlidesParseError(f"第{lineno}行: topic 標題格式應為 ## <兩位數> <名稱>")
        if line.startswith("### "):
            m = RE_PAGE.match(line)
            if not m:
                raise SlidesParseError(f"第{lineno}行: 頁標題缺版型標記, 格式應為 ### [layout] 標題")
            if topic is None:
                raise SlidesParseError(f"第{lineno}行: 頁出現在任何 ## topic 之前")
            page = Page(layout=m.group(1), title=m.group(2))
            topic.pages.append(page)
            continue
        if page is None:
            continue  # topic 標題下、第一頁之前的說明文字忽略
        if line.startswith(">"):
            page.notes = (page.notes + "\n" + line[1:].strip()).strip()
            continue
        if line.startswith("- "):
            page.bullets.append(line[2:].strip())
            continue
        m = RE_IMAGE.match(line)
        if m:
            page.images.append((m.group(1), m.group(2)))
            continue
        if line.startswith("|"):
            if RE_TABLE_SEP.match(line):
                continue
            page.table.append([c.strip() for c in line.strip().strip("|").split("|")])
            continue
        if not page.title:
            m = RE_LABEL.match(line.strip())
            if m:
                page.label, page.title = m.group(1).strip(), m.group(2).strip()
            else:
                page.title = line.strip()
            continue
        page.subtitle = (page.subtitle + "\n" + line.strip()).strip()
    if not deck.title:
        raise SlidesParseError("缺少 # 簡報名稱")
    return deck


class StyleParseError(ValueError):
    pass


@dataclass
class Style:
    slide: dict
    theme: dict
    layouts: dict


RE_YAML = re.compile(r"^```yaml\s*\n(.*?)\n^```", re.S | re.M)


def parse_style(text: str) -> Style:
    m = RE_YAML.search(text)
    if not m:
        raise StyleParseError("STYLE.md 找不到 ```yaml 區塊")
    data = yaml.safe_load(m.group(1)) or {}
    for key in ("slide", "theme", "layouts"):
        if key not in data:
            raise StyleParseError(f"STYLE.md yaml 缺少 {key}")
    return Style(data["slide"], data["theme"], data["layouts"])


VALID_ROLES = {"title", "label", "subtitle", "body", "image", "left", "right", "table", "cards"}


def validate(deck: Deck, style: Style, deck_dir: Path) -> list[str]:
    errors: list[str] = []
    for key in ("bg", "fg", "font_title", "font_body"):
        if key not in style.theme:
            errors.append(f"STYLE.md theme 缺少 {key}")
    for name, elements in style.layouts.items():
        for idx, el in enumerate(elements, 1):
            role = el.get("role")
            if role not in VALID_ROLES:
                errors.append(f"STYLE.md 版型 {name} 元素 {idx} role 不合法: {role}")
            box = el.get("box")
            if not (isinstance(box, list) and len(box) == 4 and all(isinstance(v, (int, float)) for v in box)):
                errors.append(f"STYLE.md 版型 {name} 元素 {idx} box 必須是 4 個數字")
    for topic, i, page in deck.pages():
        where = f"{topic.folder} 第{i}頁 [{page.layout}] {page.title}"
        if page.layout not in style.layouts:
            errors.append(f"{where}: 版型未在 STYLE.md 定義")
        for _desc, path in page.images:
            if path == TODO:
                continue
            full = deck_dir / path
            if not full.is_file():
                errors.append(f"{where}: 圖片不存在 {path}")
            elif full.suffix.lower() == SVG_SUFFIX:
                errors.extend(f"{where}: {e}" for e in _svg_problems(full))
    return errors


def _svg_problems(svg_path: Path) -> list[str]:
    try:
        root = ET.parse(str(svg_path)).getroot()
    except ET.ParseError as e:
        return [f"SVG 解析失敗 {svg_path.name}: {e}"]
    unsupported = _collect_unsupported_visuals(root)
    if unsupported:
        return [f"SVG 含不支援元素 {svg_path.name}: {'; '.join(unsupported[:8])}"]
    return []


PLACEHOLDER_GRAY = RGBColor(0xD1, 0xD5, 0xDB)


def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str.lstrip("#"))


def _textbox(slide, box, lines, *, font, size, color, bold=False, first_bold=False):
    x, y, w, h = (Inches(v) for v in box)
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold or (first_bold and idx == 0)
    return tb


def _picture_fit(slide, box, img_path: Path):
    x, y, w, h = box
    with Image.open(img_path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    slide.shapes.add_picture(str(img_path), Inches(px), Inches(py), Inches(pw), Inches(ph))


def _svg_size_px(root: ET.Element) -> tuple[float, float]:
    vb = root.get("viewBox")
    if vb:
        _, _, vw, vh = (float(v) for v in vb.replace(",", " ").split())
        return vw, vh
    return float(re.sub(r"[a-z%]+$", "", root.get("width"))), float(re.sub(r"[a-z%]+$", "", root.get("height")))


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _svg_group(slide, box, svg_path: Path):
    """SVG → 一個原生圖案 group, 等比縮放置中塞進 box; 群組內每個元素在 PowerPoint 裡都可個別編輯"""
    root = ET.parse(str(svg_path)).getroot()
    vw, vh = _svg_size_px(root)
    bx, by, bw, bh = box
    scale = min(bw / (vw / 96), bh / (vh / 96))
    gw, gh = (vw / 96) * scale, (vh / 96) * scale
    gx, gy = bx + (bw - gw) / 2, by + (bh - gh) / 2
    sp_tree = slide.shapes._spTree
    existing = [int(m) for m in re.findall(r'<p:cNvPr id="(\d+)"', etree.tostring(sp_tree).decode())]
    ctx = ConvertContext(defs=collect_defs(root), id_counter=max(existing, default=1) + 1)
    frags = []
    for child in root:
        if child.tag.replace(f"{{{SVG_NS}}}", "") == "defs":
            continue
        res = convert_element(child, ctx)
        if res:
            frags.append(res.xml)
    gid = ctx.next_id()
    grp_xml = (
        f'<p:grpSp xmlns:a="{A_NS}" xmlns:r="{R_NS}" xmlns:p="{P_NS}">'
        f'<p:nvGrpSpPr><p:cNvPr id="{gid}" name="diagram {svg_path.stem}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        f'<p:grpSpPr><a:xfrm><a:off x="{int(Inches(gx))}" y="{int(Inches(gy))}"/><a:ext cx="{int(Inches(gw))}" cy="{int(Inches(gh))}"/>'
        f'<a:chOff x="0" y="0"/><a:chExt cx="{int(vw * EMU_PER_PX)}" cy="{int(vh * EMU_PER_PX)}"/></a:xfrm></p:grpSpPr>'
        f'{"".join(frags)}</p:grpSp>'
    )
    sp_tree.append(etree.fromstring(grp_xml))


def _placeholder(slide, box, desc: str, theme: dict):
    x, y, w, h = (Inches(v) for v in box)
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = PLACEHOLDER_GRAY
    shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"[待補圖] {desc}"
    run.font.name = theme["font_body"]
    run.font.size = Pt(14)
    run.font.color.rgb = _rgb(theme["fg"])


def _column(page: Page, idx: int) -> list[str]:
    return [row[idx] for row in page.table if len(row) > idx]


def _table(slide, box, rows: list[list[str]], *, font, size, fg, bg, header_bg, header_fg):
    """把 page.table 畫成真正的 pptx 表格; 第一列是表頭 (加粗、header_bg 底), 其餘列 bg 底"""
    if not rows:
        return
    x, y, w, h = (Inches(v) for v in box)
    ncols = max(len(r) for r in rows)
    frame = slide.shapes.add_table(len(rows), ncols, x, y, w, h)
    tbl = frame.table
    tbl.first_row = True
    tbl.horz_banding = False
    for r, row in enumerate(rows):
        is_header = r == 0
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_bg if is_header else bg
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = font
                    run.font.size = Pt(size)
                    run.font.bold = is_header
                    run.font.color.rgb = header_fg if is_header else fg
    return frame


def _cards(slide, box, rows: list[list[str]], *, font, size, fg, cols, gap, badges, badge_fg,
           size_badge, size_body, body_color, index_color):
    """把 page.table 畫成一格一張的卡片網格; 每列 = | 徽章 | 標題 | 說明 |, 編號自動產生"""
    if not rows:
        return
    x, y, w, h = box
    nrows = -(-len(rows) // cols)
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (nrows - 1)) / nrows
    for i, row in enumerate(rows):
        cx = x + (i % cols) * (cw + gap)
        cy = y + (i // cols) * (ch + gap)
        badge = row[0] if len(row) > 0 else ""
        head = row[1] if len(row) > 1 else ""
        desc = row[2] if len(row) > 2 else ""
        bh = size_badge / 72 * 2.2
        _textbox(slide, [cx, cy, cw * 0.25, bh], [f"{i + 1:02d}"],
                 font=font, size=size_badge, color=_rgb(index_color))
        if badge:
            shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                         Inches(cx + cw * 0.22), Inches(cy),
                                         Inches(min(cw * 0.55, len(badge) * size_badge / 72 * 1.5)), Inches(bh))
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(badges.get(badge, "#6B7280"))
            shp.line.fill.background()
            shp.shadow.inherit = False
            tf = shp.text_frame
            tf.word_wrap = False
            run = tf.paragraphs[0].add_run()
            run.text = badge
            run.font.name = font
            run.font.size = Pt(size_badge)
            run.font.bold = True
            run.font.color.rgb = _rgb(badge_fg)
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        _textbox(slide, [cx, cy + bh + 0.08, cw, size / 72 * 1.8], [head],
                 font=font, size=size, color=_rgb(fg), bold=True)
        _textbox(slide, [cx, cy + bh + size / 72 * 1.8 + 0.14, cw, ch - bh - size / 72 * 1.8 - 0.14],
                 [desc], font=font, size=size_body, color=_rgb(body_color))


def render(deck: Deck, style: Style, deck_dir: Path) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(style.slide["w"])
    prs.slide_height = Inches(style.slide["h"])
    blank = prs.slide_layouts[6]
    theme = style.theme
    for _topic, _i, page in deck.pages():
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(theme["bg"])
        elements = style.layouts[page.layout]
        has_subtitle_el = any(el["role"] == "subtitle" for el in elements)
        img_i = 0
        for el in elements:
            role, box = el["role"], el["box"]
            size = el.get("size", 18)
            bold = el.get("bold", False)
            color = _rgb(el.get("color", theme["fg"]))
            if role == "title":
                _textbox(slide, box, [page.title], font=theme["font_title"], size=size, color=color, bold=bold)
            elif role == "label":
                _textbox(slide, box, [page.label], font=theme["font_title"], size=size, color=color, bold=bold)
            elif role == "subtitle":
                _textbox(slide, box, page.subtitle.splitlines(), font=theme["font_body"], size=size, color=color, bold=bold)
            elif role == "body":
                lines = page.bullets or ([] if has_subtitle_el else page.subtitle.splitlines())
                _textbox(slide, box, lines, font=theme["font_body"], size=size, color=color, bold=bold)
            elif role == "image":
                if img_i < len(page.images):
                    desc, path = page.images[img_i]
                    img_i += 1
                    if path == TODO:
                        _placeholder(slide, box, desc, theme)
                    elif path.lower().endswith(SVG_SUFFIX):
                        _svg_group(slide, box, deck_dir / path)
                    else:
                        _picture_fit(slide, box, deck_dir / path)
            elif role in ("left", "right"):
                col = 0 if role == "left" else 1
                _textbox(slide, box, _column(page, col), font=theme["font_body"], size=size, color=color, bold=bold, first_bold=True)
            elif role == "cards":
                _cards(slide, box, page.table,
                       font=theme["font_body"], size=size, fg=el.get("color", theme["fg"]),
                       cols=el.get("cols", 3), gap=el.get("gap", 0.3),
                       badges=el.get("badges", {}), badge_fg=el.get("badge_fg", theme["bg"]),
                       size_badge=el.get("size_badge", 11), size_body=el.get("size_body", 12),
                       body_color=el.get("body_color", theme["fg"]),
                       index_color=el.get("index_color", theme.get("accent", theme["fg"])))
            elif role == "table":
                _table(slide, box, page.table, font=theme["font_body"], size=size, fg=color, bg=_rgb(theme["bg"]),
                       header_bg=_rgb(el.get("header_bg", theme.get("accent", theme["fg"]))),
                       header_fg=_rgb(el.get("header_fg", theme["bg"])))
            else:
                raise StyleParseError(f"版型 {page.layout} 有未知 role: {role}")
        if page.notes:
            slide.notes_slide.notes_text_frame.text = page.notes
    return prs


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip() or "deck"


def images_todo(deck: Deck) -> str:
    rows = [
        "# 待補圖清單",
        "",
        "由 SLIDES.md 中路徑為 TODO 的圖片產生。補圖後把 SLIDES.md 的 TODO 換成實際路徑, 重跑 `--images-todo` 更新本檔。",
        "",
        "| 頁 | 描述 | 建議檔名 |",
        "|---|---|---|",
    ]
    count = 0
    for topic, i, page in deck.pages():
        for k, (desc, path) in enumerate(page.images, 1):
            if path == TODO:
                count += 1
                rows.append(f"| {topic.folder} 第{i}頁 {page.title} | {desc} | {topic.folder}/p{i:02d}_img{k}.png |")
    if count == 0:
        rows.append("| (無) | | |")
    return "\n".join(rows) + "\n"


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck_dir", help="工作資料夾, 內含 SLIDES.md 與 STYLE.md")
    ap.add_argument("--images-todo", action="store_true", help="只產 IMAGES_TODO.md, 不 build")
    args = ap.parse_args(argv)
    try:
        deck_dir = Path(args.deck_dir).resolve()
        deck = parse_slides((deck_dir / "SLIDES.md").read_text(encoding="utf-8"))
        if args.images_todo:
            (deck_dir / "IMAGES_TODO.md").write_text(images_todo(deck), encoding="utf-8")
            print(f"已寫入 {deck_dir / 'IMAGES_TODO.md'}")
            return 0
        style = parse_style((deck_dir / "STYLE.md").read_text(encoding="utf-8"))
        errors = validate(deck, style, deck_dir)
        if errors:
            print("build 失敗:")
            for e in errors:
                print(" -", e)
            return 1
        prs = render(deck, style, deck_dir)
        out = deck_dir / "output" / f"{safe_filename(deck.title)}.pptx"
        out.parent.mkdir(exist_ok=True)
        prs.save(str(out))
        print(f"已輸出 {out}")
        return 0
    except (SlidesParseError, StyleParseError, FileNotFoundError, PermissionError, KeyError) as e:
        print(f"build 失敗: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
